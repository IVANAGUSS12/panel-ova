import os
from datetime import datetime, date
import uuid

from django.core.management.base import BaseCommand, CommandError
from openpyxl import load_workbook

from core.models import Patient


class Command(BaseCommand):
    help = "Importa pacientes desde el archivo OV_AUTORIZACIONES_LOG.xlsx (hoja LOG) al modelo Patient."

    def add_arguments(self, parser):
        parser.add_argument(
            "excel_path",
            type=str,
            help="Ruta al archivo OV_AUTORIZACIONES_LOG.xlsx",
        )

    def handle(self, *args, **options):
        excel_path = options["excel_path"]

        if not os.path.exists(excel_path):
            raise CommandError(f"El archivo no existe: {excel_path}")

        wb = load_workbook(excel_path, data_only=True)
        if "LOG" not in wb.sheetnames:
            raise CommandError("La hoja 'LOG' no existe en el archivo.")

        ws = wb["LOG"]
        rows = list(ws.iter_rows(min_row=1, values_only=True))
        if not rows:
            self.stdout.write(self.style.WARNING("La hoja LOG está vacía."))
            return

        headers = [str(h).strip() if h is not None else "" for h in rows[0]]
        data_rows = rows[1:]

        # Helper para obtener índice de columna por nombre
        def idx(col_name):
            try:
                return headers.index(col_name)
            except ValueError:
                return None

        idx_timestamp = idx("timestamp")
        idx_nombre = idx("nombre")
        idx_dni = idx("dni")
        idx_email = idx("email")
        idx_telefono = idx("telefono")
        idx_cobertura = idx("cobertura")
        idx_medico = idx("medico")
        idx_observaciones = idx("observaciones")
        idx_folder_url = idx("folder_url")
        idx_subcarpeta_url = idx("subcarpeta_url")
        idx_orden_nombre = idx("orden_nombre")
        idx_orden_url = idx("orden_url")
        idx_dni_nombre = idx("dni_nombre")
        idx_dni_url = idx("dni_url")
        idx_credencial_nombre = idx("credencial_nombre")
        idx_credencial_url = idx("credencial_url")
        idx_fecha_cx = idx("fecha_cx")
        idx_estado = idx("estado")
        idx_orden_materiales_id = idx("orden_materiales_id")
        idx_presupuesto = idx("presupuesto")
        idx_reprogramacion = idx("reprogramacion")
        idx_obs_reprog = idx("observaciones_reprogramacion")
        idx_obs_prov = idx("obs_proveedores")

        created = 0
        updated = 0
        skipped = 0

        def clean_str(v):
            if v is None:
                return ""
            v = str(v).strip()
            return v

        def clean_dni(v):
            if v is None:
                return ""
            try:
                as_float = float(v)
                if as_float.is_integer():
                    return str(int(as_float))
                return str(v)
            except (TypeError, ValueError):
                return str(v).strip()

        def clean_phone(v):
            if v is None:
                return ""
            try:
                as_float = float(v)
                if as_float.is_integer():
                    return str(int(as_float))
                return str(v)
            except (TypeError, ValueError):
                return str(v).strip()

        def parse_date(v):
            if v is None or v == "":
                return None
            if isinstance(v, (datetime, date)):
                return v.date() if isinstance(v, datetime) else v
            for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
                try:
                    return datetime.strptime(str(v), fmt).date()
                except ValueError:
                    continue
            return None

        def parse_datetime(v):
            if v is None or v == "":
                return None
            if isinstance(v, datetime):
                return v
            if isinstance(v, date):
                return datetime(v.year, v.month, v.day)
            for fmt in ("%Y-%m-%d %H:%M:%S", "%d/%m/%Y %H:%M:%S"):
                try:
                    return datetime.strptime(str(v), fmt)
                except ValueError:
                    continue
            return None

        def map_status(raw):
            if raw is None:
                return Patient.STATUS_PENDIENTE
            s = str(raw).strip().lower()
            if s == "autorizado":
                return Patient.STATUS_AUTORIZADO
            if s == "solicitado":
                return Patient.STATUS_SOLICITADO
            if s == "pendiente":
                return Patient.STATUS_PENDIENTE
            if s == "autorizado material pendiente":
                return Patient.STATUS_MATERIAL_PENDIENTE
            if "rechazado" in s:
                return Patient.STATUS_RECHAZO
            return Patient.STATUS_PENDIENTE

        for row in data_rows:
            # Si no hay nombre, salto la fila
            nombre = clean_str(row[idx_nombre]) if idx_nombre is not None else ""
            if not nombre:
                skipped += 1
                continue

            dni = clean_dni(row[idx_dni]) if idx_dni is not None else ""
            email = clean_str(row[idx_email]) if idx_email is not None else ""
            telefono = clean_phone(row[idx_telefono]) if idx_telefono is not None else ""
            cobertura = clean_str(row[idx_cobertura]) if idx_cobertura is not None else ""
            medico = clean_str(row[idx_medico]) if idx_medico is not None else ""

            fecha_cx = parse_date(row[idx_fecha_cx]) if idx_fecha_cx is not None else None
            estado_raw = row[idx_estado] if idx_estado is not None else None
            timestamp = parse_datetime(row[idx_timestamp]) if idx_timestamp is not None else None

            observaciones = clean_str(row[idx_observaciones]) if idx_observaciones is not None else ""
            folder_url = clean_str(row[idx_folder_url]) if idx_folder_url is not None else ""
            subcarpeta_url = clean_str(row[idx_subcarpeta_url]) if idx_subcarpeta_url is not None else ""
            orden_nombre = clean_str(row[idx_orden_nombre]) if idx_orden_nombre is not None else ""
            orden_url = clean_str(row[idx_orden_url]) if idx_orden_url is not None else ""
            dni_nombre = clean_str(row[idx_dni_nombre]) if idx_dni_nombre is not None else ""
            dni_url = clean_str(row[idx_dni_url]) if idx_dni_url is not None else ""
            credencial_nombre = clean_str(row[idx_credencial_nombre]) if idx_credencial_nombre is not None else ""
            credencial_url = clean_str(row[idx_credencial_url]) if idx_credencial_url is not None else ""
            orden_materiales_id = clean_str(row[idx_orden_materiales_id]) if idx_orden_materiales_id is not None else ""
            presupuesto = clean_str(row[idx_presupuesto]) if idx_presupuesto is not None else ""
            reprogramacion_dt = parse_datetime(row[idx_reprogramacion]) if idx_reprogramacion is not None else None
            obs_reprog = clean_str(row[idx_obs_reprog]) if idx_obs_reprog is not None else ""
            obs_prov = clean_str(row[idx_obs_prov]) if idx_obs_prov is not None else ""

            # Servicio: todos son de Traumatología
            service = "TRAUMATOLOGIA"

            status = map_status(estado_raw)

            # Para evitar duplicados: mismo DNI + fecha_cx + cobertura + médico
            qs = Patient.objects.filter(
                dni=dni,
                planned_date=fecha_cx,
                coverage=cobertura,
                doctor=medico,
            )
            if qs.exists():
                patient = qs.first()
                created_flag = False
            else:
                patient = Patient(
                    full_name=nombre,
                    dni=dni,
                    phone=telefono or None,
                    email=email or None,
                    coverage=cobertura,
                    doctor=medico,
                    service=service,
                    planned_date=fecha_cx or date.today(),
                )
                created_flag = True

            patient.status = status

            # Observaciones externas (lo que ven los usuarios)
            external_parts = []
            if observaciones:
                external_parts.append(observaciones)
            if obs_reprog:
                external_parts.append(f"Reprogramación: {obs_reprog}")
            if obs_prov:
                external_parts.append(f"Obs. proveedores: {obs_prov}")
            patient.external_observations = "\n".join(p for p in external_parts if p)

            # Observaciones internas (info técnica / drive / etc.)
            internal_parts = []
            internal_parts.append("Importado desde LOG OVA.")
            if timestamp:
                internal_parts.append(f"Timestamp original: {timestamp.isoformat(sep=' ')}")
            if folder_url:
                internal_parts.append(f"Carpeta Drive: {folder_url}")
            if subcarpeta_url:
                internal_parts.append(f"Subcarpeta: {subcarpeta_url}")
            if orden_nombre or orden_url:
                internal_parts.append(f"Orden: {orden_nombre} ({orden_url})")
            if dni_nombre or dni_url:
                internal_parts.append(f"DNI: {dni_nombre} ({dni_url})")
            if credencial_nombre or credencial_url:
                internal_parts.append(f"Credencial: {credencial_nombre} ({credencial_url})")
            if orden_materiales_id:
                internal_parts.append(f"Orden materiales ID: {orden_materiales_id}")
            if presupuesto:
                internal_parts.append(f"Presupuesto: {presupuesto}")

            patient.internal_observations = "\n".join(p for p in internal_parts if p)

            # Reprogramación
            patient.last_reprogram_date = reprogramacion_dt
            if obs_reprog:
                patient.last_reprogram_reason = obs_reprog

            # Aseguramos un tracking_id único para los pacientes nuevos
            if created_flag and not patient.tracking_id:
                patient.tracking_id = f"OVALOG_{uuid.uuid4().hex[:10].upper()}"

            patient.save()

            if created_flag:
                created += 1
            else:
                updated += 1

        self.stdout.write(self.style.SUCCESS(
            f"Importación finalizada. Creados: {created}, Actualizados: {updated}, Saltados: {skipped}."
        ))
