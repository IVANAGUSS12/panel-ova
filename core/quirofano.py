from __future__ import annotations

from collections import defaultdict
from datetime import date, datetime, time, timedelta
from html import unescape
from html.parser import HTMLParser
from pathlib import Path
import re
import shutil

import openpyxl
from django.conf import settings
from django.db import transaction
from django.utils import timezone

from .models import Patient, QuirofanoEntry, QuirofanoSnapshot
from .report_downloader import download_qr_reports
from .text_utils import normalize_text as _shared_normalize_text


REPORT_ARTIFACT_DIRNAME = 'quirofano_reports'

SEDE_FILE_PREFIXES = {
    Patient.SEDE_SAAVEDRA: 'saavedra',
    Patient.SEDE_POMBO: 'pombo',
    Patient.SEDE_LAS_HERAS: 'las_heras',
}

SPECIALTY_ALIASES = {
    'TRAUMATOLOGIA': ['traumatologia', 'trauma', 'traumato', 'ortopedia', 'traumatolog'],
    'HEMODINAMIA': ['hemodinamia', 'hemodin'],
    'UROLOGIA': ['urologia', 'cirugia urologica', 'urologica', 'urolog'],
    'CIRUGIA GENERAL': ['cirugia general', 'cir general', 'cirugia gral', 'cir gral'],
    'CIRUGIA CABEZA Y CUELLO': ['cirugia cabeza y cuello', 'cabeza y cuello', 'cabeza cuello'],
    'CIRUGIA TORACICA': ['cirugia toracica', 'toracica', 'cir toracica', 'torax'],
    'CIRUGIA PLASTICA': ['cirugia plastica', 'plastica', 'cir plastica', 'estetica'],
    'OTORRINOLARINGOLOGIA': ['otorrinolaringologia', 'otorrino', 'orl', 'otorrinolaring'],
    'NEUROCIRUGIA': ['neurocirugia', 'neuro cirugia', 'neurocir', 'neuro'],
    'FLEBOLOGIA': ['flebologia', 'flebologa', 'flebologia'],
    'CARDIOLOGIA': ['cardiologia', 'cardiol', 'cardio'],
    'CIRUGIA VASCULAR': ['cirugia vascular', 'vascular', 'vasc'],
    'CIRUGIA CARDIOVASCULAR': ['cirugia cardiovascular', 'cardiovascular', 'cardio vasc'],
    'GINECOLOGIA': ['ginecologia', 'ginecolog', 'gine', 'ginecologia y obstetricia', 'ginecologia o/y obstetricia', 'gineco obstetricia', 'gineco obstetric'],
    'OFTALMOLOGIA': ['oftalmologia', 'oftalm', 'ojos'],
    'PEDIATRIA': ['pediatria', 'pediatr', 'pedia'],
}

DEFAULT_PANEL_SERVICES = {
    'TRAUMATOLOGIA',
    'HEMODINAMIA',
    'UROLOGIA',
    'CIRUGIA GENERAL',
    'CIRUGIA CABEZA Y CUELLO',
    'CIRUGIA TORACICA',
    'CIRUGIA PLASTICA',
    'OTORRINOLARINGOLOGIA',
    'NEUROCIRUGIA',
    'FLEBOLOGIA',
    'OFTALMOLOGIA',
    'GINECOLOGIA',
    'PEDIATRIA',
}


def normalize_text(value: str | None) -> str:
    return _shared_normalize_text(value, collapse_whitespace=True)


class _HTMLTableParser(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.rows: list[list[str]] = []
        self._current_row: list[str] | None = None
        self._current_cell: list[str] | None = None

    def handle_starttag(self, tag: str, attrs) -> None:
        if tag == 'tr':
            self._current_row = []
        elif tag in {'td', 'th'}:
            self._current_cell = []

    def handle_data(self, data: str) -> None:
        if self._current_cell is not None:
            self._current_cell.append(data)

    def handle_endtag(self, tag: str) -> None:
        if tag in {'td', 'th'} and self._current_row is not None and self._current_cell is not None:
            value = unescape(''.join(self._current_cell)).replace('\xa0', ' ')
            value = re.sub(r'\s+', ' ', value).strip()
            self._current_row.append(value)
            self._current_cell = None
        elif tag == 'tr' and self._current_row is not None:
            if any(cell.strip() for cell in self._current_row):
                self.rows.append(self._current_row)
            self._current_row = None


def normalize_dni(value: str | int | None) -> str:
    return re.sub(r'\D+', '', str(value or ''))


def add_one_month(base_date: date) -> date:
    # Usamos una ventana móvil de 31 días para que 31/03 pase a 01/05
    # en vez de quedar truncado al 30/04 por el largo del mes.
    return base_date + timedelta(days=31)


def default_quirofano_window(today: date | None = None) -> tuple[date, date]:
    today = today or date.today()
    return today, add_one_month(today)


def canonicalize_service(raw_value: str | None) -> str:
    normalized = normalize_text(raw_value)
    if not normalized:
        return ''

    for canonical, aliases in SPECIALTY_ALIASES.items():
        for alias in aliases:
            if normalized == alias or alias in normalized or normalized in alias:
                return canonical
    return ''


def _safe_filename(value: str) -> str:
    cleaned = re.sub(r'[^A-Za-z0-9._-]+', '_', value.strip())
    cleaned = cleaned.strip('._')
    return cleaned or 'reporte.xlsx'


def _read_report_rows(report_path: Path) -> list[list[str]]:
    signature = report_path.read_bytes()[:16]
    if signature.startswith(b'PK'):
        workbook = openpyxl.load_workbook(report_path, read_only=True, data_only=True)
        worksheet = workbook.active
        rows = [list(row) for row in worksheet.iter_rows(values_only=True)]
        workbook.close()
        return rows

    text = report_path.read_text(encoding='utf-8', errors='replace')
    parser = _HTMLTableParser()
    parser.feed(text)
    return parser.rows


def _report_headers_and_data(rows: list[list[str]]) -> tuple[list[str], list[list[str]]]:
    if not rows:
        return [], []

    # Formato HTML del sitio: fila 1 resumen, fila 2 títulos, fila 3+ datos.
    if len(rows) >= 2:
        header_row = rows[1]
        data_rows = rows[2:]
    else:
        header_row = rows[0]
        data_rows = rows[1:]
    return header_row, data_rows


def _build_header_index_map(headers: list[str]) -> dict[str, int]:
    header_map: dict[str, int] = {}
    for index, header in enumerate(headers):
        normalized_header = normalize_text(header)
        if normalized_header:
            header_map[normalized_header] = index
    return header_map


def _row_value(row: list, header_map: dict[str, int], *header_names: str, fallback_index: int | None = None):
    for header_name in header_names:
        index = header_map.get(normalize_text(header_name))
        if index is not None and index < len(row):
            return row[index]
    if fallback_index is not None and fallback_index < len(row):
        return row[fallback_index]
    return None


def create_artifact_directory() -> Path:
    timestamp = timezone.localtime().strftime('%Y%m%d_%H%M%S')
    artifact_dir = Path(settings.MEDIA_ROOT) / REPORT_ARTIFACT_DIRNAME / timestamp
    artifact_dir.mkdir(parents=True, exist_ok=True)
    return artifact_dir


def _relative_media_path(path: Path) -> str:
    return path.relative_to(Path(settings.MEDIA_ROOT)).as_posix()


def persist_downloaded_reports(file_paths_by_sede: dict[str, Path], artifact_dir: Path) -> dict[str, Path]:
    saved_paths: dict[str, Path] = {}
    for sede, source_path in file_paths_by_sede.items():
        prefix = SEDE_FILE_PREFIXES.get(sede, sede.lower())
        suffix = source_path.suffix or '.xlsx'
        target_path = artifact_dir / _safe_filename(f'{prefix}_diario{suffix}')
        shutil.copy2(source_path, target_path)
        saved_paths[sede] = target_path
    return saved_paths


def build_consolidated_workbook(file_paths_by_sede: dict[str, Path], artifact_dir: Path) -> Path | None:
    if not file_paths_by_sede:
        return None

    output_path = artifact_dir / 'quirofano_consolidado.xlsx'
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.title = 'Quirofano consolidado'
    worksheet.append(['Quirofano consolidado'])

    headers = None
    for sede, report_path in file_paths_by_sede.items():
        source_rows = _read_report_rows(report_path)
        source_headers, data_rows = _report_headers_and_data(source_rows)
        if not source_headers:
            continue

        if headers is None:
            headers = list(source_headers) + ['SEDE_PANEL', 'ARCHIVO_ORIGEN']
            worksheet.append(headers)

        for row in data_rows:
            if not any(value not in (None, '') for value in row):
                continue
            worksheet.append(list(row) + [sede, report_path.name])

    workbook.save(output_path)
    workbook.close()
    return output_path if headers is not None else None


def build_source_metadata(file_paths_by_sede: dict[str, Path], merged_path: Path | None) -> list[dict]:
    metadata: list[dict] = []
    for sede, report_path in file_paths_by_sede.items():
        metadata.append({
            'kind': 'source',
            'sede': sede,
            'name': report_path.name,
            'stored_path': _relative_media_path(report_path),
        })

    if merged_path is not None:
        metadata.append({
            'kind': 'merged',
            'name': merged_path.name,
            'stored_path': _relative_media_path(merged_path),
        })
    return metadata


def _coerce_date(value) -> date | None:
    if value is None or value == '':
        return None
    if isinstance(value, date) and not isinstance(value, datetime):
        return value
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, str):
        cleaned = value.strip()
        for fmt in ('%d/%m/%Y', '%Y-%m-%d', '%d-%m-%Y'):
            try:
                return datetime.strptime(cleaned, fmt).date()
            except ValueError:
                continue
    return None


def _coerce_time(value) -> time | None:
    if value is None or value == '':
        return None
    if isinstance(value, datetime):
        return value.time().replace(second=0, microsecond=0)
    if isinstance(value, time):
        return value.replace(second=0, microsecond=0)
    if isinstance(value, str):
        cleaned = value.strip()
        for fmt in ('%H:%M', '%H:%M:%S'):
            try:
                return datetime.strptime(cleaned, fmt).time().replace(second=0, microsecond=0)
            except ValueError:
                continue
    return None


def _allowed_panel_services() -> set[str]:
    allowed = set(DEFAULT_PANEL_SERVICES)
    for service in Patient.objects.exclude(service__isnull=True).exclude(service__exact='').values_list('service', flat=True).distinct():
        canonical = canonicalize_service(service)
        if canonical:
            allowed.add(canonical)
    return allowed


def _is_allowed_origin(raw_value: str | None) -> bool:
    origin = normalize_text(raw_value)
    return 'admisi' in origin


def parse_report_workbook(report_path: Path, sede: str, date_from: date, date_to: date, allowed_services: set[str]) -> list[dict]:
    source_rows = _read_report_rows(report_path)
    headers, data_rows = _report_headers_and_data(source_rows)
    header_map = _build_header_index_map(headers)
    rows: list[dict] = []

    for row_index, row in enumerate(data_rows, start=3):
        surgery_date = _coerce_date(_row_value(row, header_map, 'Fecha de Cirugia', fallback_index=1))
        if not surgery_date or surgery_date < date_from or surgery_date > date_to:
            continue

        specialty_raw = str(_row_value(row, header_map, 'Especialidad', fallback_index=26) or '').strip()
        canonical_service = canonicalize_service(specialty_raw)
        if not canonical_service or canonical_service not in allowed_services:
            continue

        origin_raw = str(_row_value(row, header_map, 'Origen', fallback_index=27) or '').strip()
        if not _is_allowed_origin(origin_raw):
            continue

        patient_name = str(_row_value(row, header_map, 'Paciente', fallback_index=7) or '').strip()
        dni = normalize_dni(_row_value(row, header_map, 'Dni', fallback_index=10))
        if not patient_name and not dni:
            continue

        report_phone = str(_row_value(row, header_map, 'Telefono', fallback_index=11) or '').strip()
        destination_service = str(_row_value(row, header_map, 'Servicio de Destino', fallback_index=25) or '').strip()

        rows.append({
            'sede': sede,
            'surgery_date': surgery_date,
            'surgery_time': _coerce_time(_row_value(row, header_map, 'Hora Programada', fallback_index=2)),
            'patient_name': patient_name,
            'patient_name_norm': normalize_text(patient_name),
            'coverage': str(_row_value(row, header_map, 'Obra Social', fallback_index=9) or '').strip(),
            'dni': dni,
            'report_phone': report_phone,
            'doctor': str(_row_value(row, header_map, 'Cirujano', fallback_index=15) or '').strip(),
            'doctor_norm': normalize_text(_row_value(row, header_map, 'Cirujano', fallback_index=15)),
            'destination_service': destination_service,
            'specialty_raw': specialty_raw,
            'canonical_service': canonical_service,
            'origin': origin_raw,
            'source_file': report_path.name,
            'source_row_number': row_index,
        })

    return rows


def _build_match_key(row: dict) -> str:
    if row.get('dni'):
        return f"dni:{row['dni']}"
    if row.get('patient_name_norm'):
        return f"name:{row['patient_name_norm']}"
    return f"row:{row.get('source_file', '')}:{row.get('source_row_number', 0)}"


def _row_signature(row: dict) -> tuple:
    return (
        row.get('sede') or '',
        row.get('surgery_date'),
        row.get('surgery_time'),
        row.get('coverage') or '',
        row.get('doctor_norm') or '',
        row.get('canonical_service') or '',
        normalize_text(row.get('origin') or ''),
    )


def _build_previous_lookup(previous_snapshot: QuirofanoSnapshot | None) -> dict[str, list[dict]]:
    previous_lookup: dict[str, list[dict]] = defaultdict(list)
    if previous_snapshot is None:
        return previous_lookup

    previous_entries = previous_snapshot.entries.exclude(change_type=QuirofanoEntry.CHANGE_REMOVED)
    for entry in previous_entries:
        previous_lookup[entry.match_key].append({
            'entry': entry,
            'used': False,
            'signature': (
                entry.sede,
                entry.surgery_date,
                entry.surgery_time,
                entry.coverage,
                entry.doctor_norm,
                entry.canonical_service,
                normalize_text(entry.origin),
            ),
        })
    return previous_lookup


def _pick_previous_candidate(row: dict, candidates: list[dict]) -> dict | None:
    available = [candidate for candidate in candidates if not candidate['used']]
    if not available:
        return None

    target_signature = _row_signature(row)
    for candidate in available:
        if candidate['signature'] == target_signature:
            return candidate

    def _score(candidate: dict) -> tuple:
        entry = candidate['entry']
        return (
            1 if entry.canonical_service == row.get('canonical_service') else 0,
            1 if entry.sede == row.get('sede') else 0,
            1 if entry.surgery_date == row.get('surgery_date') else 0,
            1 if entry.surgery_time == row.get('surgery_time') else 0,
        )

    available.sort(key=_score, reverse=True)
    return available[0]


def _match_app_patient(row: dict, patients_by_dni: dict[str, Patient], patients_by_name: dict[str, list[Patient]]) -> Patient | None:
    surgery_date: date | None = row.get('surgery_date')
    dni = row.get('dni') or ''
    if dni and dni in patients_by_dni:
        candidate = patients_by_dni[dni]
        # Solo matchear si la fecha de la app está dentro de 20 días de la fecha de quirófano
        if surgery_date and candidate.planned_date:
            if abs((candidate.planned_date - surgery_date).days) > 20:
                return None
        return candidate

    name_norm = row.get('patient_name_norm') or ''
    if not name_norm:
        return None

    candidates = patients_by_name.get(name_norm, [])
    if not candidates:
        return None

    for candidate in candidates:
        if candidate.planned_date == surgery_date:
            return candidate
    return candidates[0]


def _build_app_match_key(patient: Patient) -> str:
    dni = normalize_dni(patient.dni)
    if dni:
        return f'dni:{dni}'
    name_norm = patient.full_name_norm or normalize_text(patient.full_name)
    if name_norm:
        return f'name:{name_norm}'
    return f'app:{patient.pk}'


@transaction.atomic
def import_daily_quirofano_snapshot(*, user, username: str | None = None, password: str | None = None) -> QuirofanoSnapshot:
    date_from, date_to = default_quirofano_window()
    artifact_dir = create_artifact_directory()
    temp_download_dir = artifact_dir / 'raw_downloads'

    downloaded_paths = download_qr_reports(
        date_from=date_from,
        date_to=date_to,
        output_dir=temp_download_dir,
        username=username,
        password=password,
    )

    stored_paths = persist_downloaded_reports(downloaded_paths, artifact_dir)
    merged_path = build_consolidated_workbook(stored_paths, artifact_dir)
    source_files = build_source_metadata(stored_paths, merged_path)
    allowed_services = _allowed_panel_services()

    current_rows: list[dict] = []
    for sede, report_path in stored_paths.items():
        current_rows.extend(parse_report_workbook(report_path, sede, date_from, date_to, allowed_services))

    previous_snapshot = QuirofanoSnapshot.objects.first()
    previous_lookup = _build_previous_lookup(previous_snapshot)

    patients = list(
        Patient.objects.filter(
            planned_date__gte=date_from,
            planned_date__lte=date_to,
        )
    )
    patients_by_dni = {normalize_dni(patient.dni): patient for patient in patients if normalize_dni(patient.dni)}
    patients_by_name: dict[str, list[Patient]] = defaultdict(list)
    for patient in patients:
        name_norm = patient.full_name_norm or normalize_text(patient.full_name)
        if name_norm:
            patients_by_name[name_norm].append(patient)

    snapshot = QuirofanoSnapshot.objects.create(
        created_by=user,
        window_start=date_from,
        window_end=date_to,
        source_files=source_files,
    )

    new_count = 0
    changed_count = 0
    unchanged_count = 0
    matched_patient_ids: set[int] = set()

    for row in current_rows:
        row['match_key'] = _build_match_key(row)
        candidate = _pick_previous_candidate(row, previous_lookup.get(row['match_key'], []))
        change_type = QuirofanoEntry.CHANGE_NEW
        previous_values = {
            'previous_sede': '',
            'previous_surgery_date': None,
            'previous_surgery_time': None,
            'previous_doctor': '',
            'previous_coverage': '',
            'previous_specialty_raw': '',
            'previous_origin': '',
        }

        if candidate is not None:
            candidate['used'] = True
            previous_entry = candidate['entry']
            previous_values = {
                'previous_sede': previous_entry.sede,
                'previous_surgery_date': previous_entry.surgery_date,
                'previous_surgery_time': previous_entry.surgery_time,
                'previous_doctor': previous_entry.doctor,
                'previous_coverage': previous_entry.coverage,
                'previous_specialty_raw': previous_entry.specialty_raw,
                'previous_origin': previous_entry.origin,
            }
            if candidate['signature'] == _row_signature(row):
                change_type = QuirofanoEntry.CHANGE_UNCHANGED
                unchanged_count += 1
            else:
                change_type = QuirofanoEntry.CHANGE_CHANGED
                changed_count += 1
        else:
            new_count += 1

        app_patient = _match_app_patient(row, patients_by_dni, patients_by_name)
        comparison_status = (
            QuirofanoEntry.COMPARISON_BOTH if app_patient is not None else QuirofanoEntry.COMPARISON_ONLY_QUIROFANO
        )
        resolved_workflow_status = app_patient.status if app_patient is not None else Patient.STATUS_PENDIENTE
        if app_patient is not None:
            matched_patient_ids.add(app_patient.id)

        QuirofanoEntry.objects.create(
            snapshot=snapshot,
            app_patient=app_patient,
            change_type=change_type,
            comparison_status=comparison_status,
            resolved_comparison_status=comparison_status,
            resolved_workflow_status=resolved_workflow_status,
            match_key=row['match_key'],
            sede=row['sede'],
            surgery_date=row['surgery_date'],
            surgery_time=row['surgery_time'],
            patient_name=row['patient_name'],
            patient_name_norm=row['patient_name_norm'],
            coverage=row['coverage'],
            dni=row['dni'],
            report_phone=row['report_phone'],
            doctor=row['doctor'],
            doctor_norm=row['doctor_norm'],
            destination_service=row['destination_service'],
            specialty_raw=row['specialty_raw'],
            canonical_service=row['canonical_service'],
            origin=row['origin'],
            source_file=row['source_file'],
            source_row_number=row['source_row_number'],
            **previous_values,
        )

    removed_count = 0
    for candidates in previous_lookup.values():
        for candidate in candidates:
            if candidate['used']:
                continue
            previous_entry = candidate['entry']
            removed_count += 1
            QuirofanoEntry.objects.create(
                snapshot=snapshot,
                app_patient=previous_entry.app_patient,
                change_type=QuirofanoEntry.CHANGE_REMOVED,
                comparison_status=previous_entry.comparison_status,
                resolved_comparison_status=previous_entry.resolved_comparison_status,
                manual_workflow_status=previous_entry.manual_workflow_status,
                resolved_workflow_status=previous_entry.resolved_workflow_status,
                match_key=previous_entry.match_key,
                sede=previous_entry.sede,
                surgery_date=previous_entry.surgery_date,
                surgery_time=previous_entry.surgery_time,
                patient_name=previous_entry.patient_name,
                patient_name_norm=previous_entry.patient_name_norm,
                coverage=previous_entry.coverage,
                dni=previous_entry.dni,
                report_phone=previous_entry.report_phone,
                doctor=previous_entry.doctor,
                doctor_norm=previous_entry.doctor_norm,
                destination_service=previous_entry.destination_service,
                specialty_raw=previous_entry.specialty_raw,
                canonical_service=previous_entry.canonical_service,
                origin=previous_entry.origin,
                source_file=previous_entry.source_file,
                source_row_number=previous_entry.source_row_number,
                previous_sede=previous_entry.sede,
                previous_surgery_date=previous_entry.surgery_date,
                previous_surgery_time=previous_entry.surgery_time,
                previous_doctor=previous_entry.doctor,
                previous_coverage=previous_entry.coverage,
                previous_specialty_raw=previous_entry.specialty_raw,
                previous_origin=previous_entry.origin,
            )

    for patient in patients:
        if patient.id in matched_patient_ids:
            continue

        canonical_service = canonicalize_service(patient.service)
        if not canonical_service or canonical_service not in allowed_services:
            continue

        patient_name_norm = patient.full_name_norm or normalize_text(patient.full_name)
        QuirofanoEntry.objects.create(
            snapshot=snapshot,
            app_patient=patient,
            change_type=QuirofanoEntry.CHANGE_UNCHANGED,
            comparison_status=QuirofanoEntry.COMPARISON_ONLY_APP,
            resolved_comparison_status=QuirofanoEntry.COMPARISON_ONLY_APP,
            resolved_workflow_status=patient.status,
            match_key=_build_app_match_key(patient),
            sede=(patient.sede or '').strip(),
            surgery_date=patient.planned_date,
            surgery_time=patient.surgery_time,
            patient_name=patient.full_name,
            patient_name_norm=patient_name_norm,
            coverage=patient.coverage or '',
            dni=normalize_dni(patient.dni),
            report_phone='',
            doctor=patient.doctor or '',
            doctor_norm=normalize_text(patient.doctor),
            destination_service='',
            specialty_raw=patient.service or '',
            canonical_service=canonical_service,
            origin='APP',
            source_file='',
            source_row_number=0,
        )

    snapshot.total_entries = len(current_rows)
    snapshot.new_count = new_count
    snapshot.changed_count = changed_count
    snapshot.unchanged_count = unchanged_count
    snapshot.removed_count = removed_count
    snapshot.save(update_fields=['total_entries', 'new_count', 'changed_count', 'unchanged_count', 'removed_count'])
    return snapshot